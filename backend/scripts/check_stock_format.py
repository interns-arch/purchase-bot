"""Check: a stock file without MRP gets the vendor and the purchase team
asked for our format, once a day; a good file and an own warehouse are not
(24 Oct 2026, Mohan Spares).

Throwaway SQLite, never the real database; WhatsApp sends are captured:

    .venv\\Scripts\\python.exe -m backend.scripts.check_stock_format
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

_TMP = Path(tempfile.mkdtemp(prefix="stock-format-"))
os.environ["DATABASE_URL"] = f"sqlite:///{(_TMP / 'check.db').as_posix()}"
os.environ["STOCK_FORMAT_TEAM_NUMBERS"] = "Seema:9289015770,Alam:9217030408,Rajkumar:9217396458"

failures = 0


def check(name: str, cond: bool) -> None:
    global failures
    print(("  PASS " if cond else "  FAIL ") + name)
    if not cond:
        failures += 1


def stock_file(name: str, headers: list[str], rows: list[tuple]) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    path = _TMP / name
    wb.save(path)
    return path


def main() -> int:
    from core.db import get_session, init_db
    from core.models import Vendor
    from core.services import inventory_import_service as svc
    from backend.app.integrations.whatsapp import outbound, registry, stock_format

    sent: list[tuple[str, str]] = []
    docs: list[str] = []
    outbound.send_reply_safe = lambda to, body: sent.append((to, body)) or True
    outbound.send_document_safe = lambda to, *a, **k: docs.append(to) or True
    outbound.in_service_window = lambda to: True

    init_db(force=True)
    no_mrp = stock_file("stock_Mohan.xlsx", ["Part No", "Part Name", "Stock Qty", "Brand"],
                        [("0910310216", "BOLT", 28, "MARUTI"), ("0920605001", "PIN", 10, "MARUTI")])
    with_mrp = stock_file("good.xlsx", ["PartNum", "Part Description", "MRP", "Current Stock"],
                          [("0910310216", "BOLT", 89, 28)])

    with get_session() as s:
        mohan = Vendor(name="Mohan Spares Parts Pvt. Ltd.", vendor_code="CT_MSP")
        good = Vendor(name="Good Motors", vendor_code="CT_GM")
        own = Vendor(name="BIJWASHAN STOCK", vendor_code="CT_BS")
        s.add_all([mohan, good, own])
        s.flush()
        mid, gid, oid = mohan.id, good.id, own.id
        registry.register_vendor_number("9711175797", mid, s)
        svc.run_import(mid, no_mrp, s)
        svc.run_import(gid, with_mrp, s)
        svc.run_import(oid, no_mrp, s)

    print("\n[1] file without MRP -> vendor and team asked")
    asked = stock_format._check_and_ask(mid)
    to = [t for t, _ in sent]
    check("asked", asked)
    check("vendor's registered number got the format", "919711175797" in to)
    check("Seema, Alam, Rajkumar got it", {"919289015770", "919217030408", "919217396458"} <= set(to))
    check("message names MRP and our format", all("MRP" in b and stock_format.FORMAT_LINE in b for _, b in sent))
    check("sample Excel sent to the vendor (chat open)", docs == ["919711175797"])
    check("team told the vendor was informed", "bata diya" in dict(sent)["919217030408"])

    print("\n[2] second file the same day -> not asked again")
    sent.clear()
    check("not asked twice", not stock_format._check_and_ask(mid) and not sent)

    print("\n[3] file with MRP -> nobody asked")
    check("good vendor not asked", not stock_format._check_and_ask(gid))

    print("\n[4] own warehouse without MRP -> nobody asked")
    check("Bijwashan not asked", not stock_format._check_and_ask(oid))

    print("\n[5] vendor with no registered number -> team told the bot could not ask")
    stock_format._asked.clear()
    with get_session() as s:
        registry.unregister("9711175797", s)
    sent.clear()
    stock_format._check_and_ask(mid)
    check("only the team messaged", {t for t, _ in sent} == {"919289015770", "919217030408", "919217396458"})
    check("team told the number is not registered", "registered nahi" in sent[0][1])

    print("\n[6] the sample Excel opens with our headers")
    import openpyxl, io

    ws = openpyxl.load_workbook(io.BytesIO(stock_format.sample_excel())).active
    check("headers", [c.value for c in ws[1]] == stock_format.FORMAT_HEADERS)

    print(f"\n{'ALL PASSED' if not failures else f'{failures} FAILED'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
